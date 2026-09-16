# ==============================================================================
# MÓDULO: reportes/exporters.py
# DESCRIPCIÓN: Exportadores profesionales para Excel (XLSX), CSV y HTML con
#              estándares institucionales y médicos de SIGEPSI.
# ==============================================================================
import csv
import io
from datetime import datetime
from django.http import HttpResponse

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExcelExporter:
    """
    Genera un libro de Excel (.xlsx) altamente estético, formal y firme,
    con membrete institucional SIGEPSI, encabezados estilizados, bordes finos,
    formatos numéricos de moneda y ancho automático de columnas.
    """

    @staticmethod
    def generar(resultado, nombre="reporte", tenant_name="Centro Psicológico"):
        if not OPENPYXL_AVAILABLE:
            # Fallback a CSV si openpyxl no estuviera disponible
            return CSVExporter.generar(resultado, nombre)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte Clínico"
        ws.views.sheetView[0].showGridLines = True

        # Colores corporativos SIGEPSI
        COLOR_PRIMARY_DARK = "0F2922"   # Verde bosque oscuro
        COLOR_PRIMARY = "19734E"        # Verde esmeralda clínico
        COLOR_ACCENT = "2EC486"         # Verde menta
        COLOR_ZEBRA = "F7FAF8"          # Alternado muy suave
        COLOR_BORDER = "D3E0D8"         # Borde sutil

        # Fuentes
        font_main_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        font_subtitle = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_meta = Font(name="Calibri", size=9, italic=True, color="475569")
        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10, color="1E293B")
        font_total = Font(name="Calibri", size=10, bold=True, color="0F2922")

        # Rellenos
        fill_main_title = PatternFill(start_color=COLOR_PRIMARY_DARK, end_color=COLOR_PRIMARY_DARK, fill_type="solid")
        fill_subtitle = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
        fill_header = PatternFill(start_color="164E3D", end_color="164E3D", fill_type="solid")
        fill_zebra = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
        fill_summary = PatternFill(start_color="E8F5EE", end_color="E8F5EE", fill_type="solid")

        # Bordes
        thin_side = Side(border_style="thin", color=COLOR_BORDER)
        double_side = Side(border_style="double", color="0F2922")
        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        border_total = Border(top=thin_side, bottom=double_side)

        col_labels = [c["label"] for c in resultado["columnas"]]
        col_keys = [c["key"] for c in resultado["columnas"]]
        num_cols = max(len(col_labels), 4)

        # Fila 1: Título Institucional
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        cell_t1 = ws.cell(row=1, column=1, value="SISTEMA INTEGRADO DE GESTIÓN PSICOLÓGICA — SIGEPSI")
        cell_t1.font = font_main_title
        cell_t1.fill = fill_main_title
        cell_t1.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        # Fila 2: Subtítulo con nombre del reporte y centro
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
        titulo_rep = f"REPORTE OFICIAL: {nombre.upper().replace('_', ' ')} — {tenant_name.upper()}"
        cell_t2 = ws.cell(row=2, column=1, value=titulo_rep)
        cell_t2.font = font_subtitle
        cell_t2.fill = fill_subtitle
        cell_t2.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 24

        # Fila 3: Metadatos y fecha de emisión
        fecha_emision = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_cols)
        cell_meta = ws.cell(
            row=3, column=1,
            value=f"Emisión: {fecha_emision} | Total de registros: {resultado['total']} | Clasificación: Documento Clínico Confidencial"
        )
        cell_meta.font = font_meta
        cell_meta.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 18

        # Fila 4: Espacio
        ws.row_dimensions[4].height = 8

        # Fila 5: Cabeceras de columna
        header_row = 5
        ws.row_dimensions[header_row].height = 24
        for col_idx, label in enumerate(col_labels, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=label)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_cell

        # Filas de datos
        current_row = 6
        for fila in resultado["datos"]:
            ws.row_dimensions[current_row].height = 20
            is_even = (current_row % 2 == 0)

            for col_idx, key in enumerate(col_keys, 1):
                val = fila.get(key, "")
                cell = ws.cell(row=current_row, column=col_idx)

                # Formateo inteligente según el tipo de campo
                if key in ["costo", "monto", "tarifa_base"] and val not in ["", None]:
                    try:
                        cell.value = float(val)
                        cell.number_format = '"Bs." #,##0.00'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    except (ValueError, TypeError):
                        cell.value = str(val)
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                elif key in ["fecha", "hora_inicio", "hora_fin", "ci", "codigo_expediente"]:
                    cell.value = str(val) if val is not None else ""
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.value = str(val) if val is not None else ""
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                cell.font = font_data
                cell.border = border_cell
                if is_even:
                    cell.fill = fill_zebra

            current_row += 1

        # Resumen al pie si existe
        resumen = resultado.get("resumen", {})
        if resumen:
            current_row += 1
            ws.row_dimensions[current_row].height = 22
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=num_cols)
            cell_res_header = ws.cell(row=current_row, column=1, value="RESUMEN EJECUTIVO Y ESTADÍSTICO")
            cell_res_header.font = Font(name="Calibri", size=10, bold=True, color="0F2922")
            cell_res_header.fill = fill_summary
            cell_res_header.alignment = Alignment(horizontal="left", vertical="center")
            cell_res_header.border = border_cell

            current_row += 1
            for k, v in resumen.items():
                ws.row_dimensions[current_row].height = 18
                label_res = k.replace("_", " ").capitalize()
                if isinstance(v, dict):
                    ws.cell(row=current_row, column=1, value=label_res).font = font_total
                    current_row += 1
                    for sk, sv in v.items():
                        ws.cell(row=current_row, column=1, value=f"  • {sk}:").font = font_data
                        ws.cell(row=current_row, column=2, value=str(sv)).font = font_data
                        current_row += 1
                else:
                    ws.cell(row=current_row, column=1, value=f"• {label_res}:").font = font_total
                    if "ingreso" in k or "costo" in k or "monto" in k:
                        try:
                            cell_val = ws.cell(row=current_row, column=2, value=float(v))
                            cell_val.number_format = '"Bs." #,##0.00'
                            cell_val.font = font_total
                        except Exception:
                            ws.cell(row=current_row, column=2, value=str(v)).font = font_total
                    else:
                        ws.cell(row=current_row, column=2, value=str(v)).font = font_total
                    current_row += 1

        # Bloque de Firmas Institucionales al pie
        current_row += 2
        ws.row_dimensions[current_row].height = 30
        col_sig1 = 1
        col_sig2 = max(num_cols - 1, 3)

        ws.cell(row=current_row, column=col_sig1, value="_______________________________________").font = font_meta
        ws.cell(row=current_row, column=col_sig2, value="_______________________________________").font = font_meta
        current_row += 1
        ws.cell(row=current_row, column=col_sig1, value="Responsable de Emisión / Administración").font = font_total
        ws.cell(row=current_row, column=col_sig2, value="Dirección Médica / Validación Clínica").font = font_total

        # Ajuste automático del ancho de columnas (Auto-fit con margen de seguridad)
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                # Omitir las primeras filas combinadas para el cálculo
                if cell.row in [1, 2, 3]:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # Generar buffer en memoria
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="reporte_{nombre}.xlsx"'
        return response


class CSVExporter:
    """Genera una respuesta HTTP con descarga CSV estructurada con BOM UTF-8."""

    @staticmethod
    def generar(resultado, nombre="reporte"):
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="reporte_{nombre}.csv"'
        response.write('\ufeff')  # BOM for Excel UTF-8 compatibility

        writer = csv.writer(response)

        # Título y metadatos en cabecera
        writer.writerow(["SISTEMA INTEGRADO DE GESTIÓN PSICOLÓGICA — SIGEPSI"])
        writer.writerow([f"REPORTE OFICIAL: {nombre.upper().replace('_', ' ')}", f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        writer.writerow([])

        # Header
        col_labels = [c["label"] for c in resultado["columnas"]]
        col_keys = [c["key"] for c in resultado["columnas"]]
        writer.writerow(col_labels)

        # Data
        for fila in resultado["datos"]:
            row = []
            for key in col_keys:
                val = fila.get(key, "")
                row.append(str(val) if val is not None else "")
            writer.writerow(row)

        # Resumen
        resumen = resultado.get("resumen", {})
        if resumen:
            writer.writerow([])
            writer.writerow(["--- RESUMEN ESTADÍSTICO ---"])
            for k, v in resumen.items():
                if isinstance(v, dict):
                    writer.writerow([k])
                    for sk, sv in v.items():
                        writer.writerow(["", sk, sv])
                else:
                    writer.writerow([k, v])

        return response


class HTMLExporter:
    """Genera una vista previa e informe HTML con presentación formal y membrete clínico."""

    @staticmethod
    def generar(resultado, titulo="Reporte SIGEPSI", tenant_name="Centro Clínico"):
        col_labels = [c["label"] for c in resultado["columnas"]]
        col_keys = [c["key"] for c in resultado["columnas"]]
        fecha_emision = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

        html = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<title>{titulo} — SIGEPSI</title>
<style>
  @page {{ size: A4 landscape; margin: 15mm; }}
  body {{
    font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, sans-serif;
    margin: 20px 30px;
    color: #12271f;
    background: #fdfefe;
    font-size: 13px;
  }}
  .membrete {{
    border-bottom: 2px solid #0f2922;
    padding-bottom: 12px;
    margin-bottom: 20px;
    display: flex;
    justify-content: space-between;
    align-items: flex-end;
  }}
  .membrete-brand {{
    display: flex;
    align-items: center;
    gap: 12px;
  }}
  .logo-badge {{
    width: 44px;
    height: 44px;
    background: #0f2922;
    color: #ffffff;
    font-size: 22px;
    font-weight: 700;
    display: flex;
    align-items: center;
    justify-content: center;
    border-radius: 8px;
  }}
  .membrete-title h1 {{
    margin: 0;
    font-size: 18px;
    color: #0f2922;
    letter-spacing: 0.5px;
  }}
  .membrete-title p {{
    margin: 3px 0 0 0;
    font-size: 12px;
    color: #19734e;
    font-weight: 600;
  }}
  .meta-box {{
    text-align: right;
    font-size: 11px;
    color: #557568;
    line-height: 1.4;
  }}
  .meta-box strong {{ color: #0f2922; }}
  .report-title-bar {{
    background: linear-gradient(135deg, #19734e 0%, #0f2922 100%);
    color: #ffffff;
    padding: 10px 16px;
    border-radius: 6px;
    display: flex;
    justify-content: space-between;
    align-items: center;
    margin-bottom: 16px;
  }}
  .report-title-bar h2 {{ margin: 0; font-size: 14px; text-transform: uppercase; letter-spacing: 0.5px; }}
  .report-title-bar span {{ font-size: 12px; background: rgba(255,255,255,0.15); padding: 3px 8px; border-radius: 4px; }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin-top: 10px;
    background: #ffffff;
    box-shadow: 0 1px 3px rgba(0,0,0,0.05);
  }}
  th {{
    background: #164e3d;
    color: #ffffff;
    padding: 9px 12px;
    text-align: left;
    font-size: 11px;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.4px;
    border: 1px solid #164e3d;
  }}
  td {{
    border: 1px solid #e2ece6;
    padding: 8px 12px;
    font-size: 12px;
    color: #1f362c;
  }}
  tr:nth-child(even) {{ background: #f8faf9; }}
  .resumen-box {{
    margin-top: 24px;
    background: #eef6f2;
    border: 1px solid #c9e2d5;
    border-left: 4px solid #19734e;
    padding: 12px 18px;
    border-radius: 6px;
  }}
  .resumen-box h3 {{ margin: 0 0 8px 0; font-size: 12px; color: #0f2922; text-transform: uppercase; }}
  .firmas-grid {{
    margin-top: 40px;
    display: flex;
    justify-content: space-between;
    padding: 0 40px;
  }}
  .firma-col {{
    text-align: center;
    width: 250px;
    border-top: 1px solid #0f2922;
    padding-top: 6px;
    font-size: 11px;
    font-weight: 600;
    color: #3b5a4e;
  }}
  .footer-legal {{
    margin-top: 30px;
    text-align: center;
    font-size: 10px;
    color: #8da89d;
    border-top: 1px solid #eef4f1;
    padding-top: 10px;
  }}
</style>
</head>
<body>
  <div class="membrete">
    <div class="membrete-brand">
      <div class="logo-badge">Ψ</div>
      <div class="membrete-title">
        <h1>SIGEPSI — PLATAFORMA CLÍNICA SAAS</h1>
        <p>{tenant_name}</p>
      </div>
    </div>
    <div class="meta-box">
      <div>Folio Oficial: <strong>REP-{datetime.now().strftime('%Y%m%d%H%M')}</strong></div>
      <div>Fecha de Emisión: <strong>{fecha_emision}</strong></div>
      <div>Validez: <strong>Documento Clínico Oficial</strong></div>
    </div>
  </div>

  <div class="report-title-bar">
    <h2>{titulo}</h2>
    <span>{resultado['total']} Registros generados</span>
  </div>

  <table>
    <thead><tr>"""

        for label in col_labels:
            html += f"<th>{label}</th>"
        html += "</tr></thead><tbody>"

        for fila in resultado["datos"]:
            html += "<tr>"
            for key in col_keys:
                val = fila.get(key, "")
                if key in ["costo", "monto", "tarifa_base"] and val not in ["", None]:
                    try:
                        val_str = f"Bs. {float(val):,.2f}"
                    except Exception:
                        val_str = str(val)
                else:
                    val_str = str(val) if val is not None else ""
                html += f"<td>{val_str}</td>"
            html += "</tr>"

        html += "</tbody></table>"

        resumen = resultado.get("resumen", {})
        if resumen:
            html += '<div class="resumen-box"><h3>Resumen Estadístico</h3><ul>'
            for k, v in resumen.items():
                if isinstance(v, dict):
                    html += f"<li><strong>{k.replace('_', ' ').title()}:</strong><ul>"
                    for sk, sv in v.items():
                        html += f"<li>{sk}: {sv}</li>"
                    html += "</ul></li>"
                else:
                    html += f"<li><strong>{k.replace('_', ' ').title()}:</strong> {v}</li>"
            html += "</ul></div>"

        html += """
  <div class="firmas-grid">
    <div class="firma-col">
      Responsable de Emisión / Administración
    </div>
    <div class="firma-col">
      Dirección Médica / Sello de Validación
    </div>
  </div>

  <div class="footer-legal">
    Este reporte ha sido generado por el Sistema Integrado de Gestión Psicológica (SIGEPSI). Información sujeta a reserva legal y confidencialidad médica.
  </div>
</body>
</html>"""
        return html
